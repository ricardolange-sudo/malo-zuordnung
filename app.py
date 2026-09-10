import io
import re
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

APP_DIR = Path.home() / ".malo_zuordnung"
APP_DIR.mkdir(exist_ok=True)
DB_PATH = APP_DIR / "malo_zuordnung.db"

st.set_page_config(page_title="MaLo-Zuordnung", page_icon="⚡", layout="wide")


@contextmanager
def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def ensure_column(table, column, column_type="TEXT"):
    with db() as conn:
        existing = [row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()]
        if column not in existing:
            conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {column_type}")


def init_db():
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS objects (
            id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, street TEXT, house_number TEXT,
            postal_code TEXT, city TEXT, address TEXT, contact TEXT, network_operator TEXT, created_at TEXT NOT NULL
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS assignments (
            id INTEGER PRIMARY KEY AUTOINCREMENT, object_id INTEGER NOT NULL, unit TEXT NOT NULL, location TEXT,
            usage_type TEXT NOT NULL, malo_id TEXT, meter_number TEXT, meter_location TEXT,
            measurement_location TEXT, tenant TEXT, status TEXT NOT NULL, checked_on TEXT, note TEXT,
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL, FOREIGN KEY(object_id) REFERENCES objects(id)
        )""")
    for column in ["street", "house_number", "postal_code", "city"]:
        ensure_column("objects", column)


def clean(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def object_address(row):
    street_line = " ".join(part for part in [clean(row.get("street", "")), clean(row.get("house_number", ""))] if part)
    city_line = " ".join(part for part in [clean(row.get("postal_code", "")), clean(row.get("city", ""))] if part)
    return ", ".join(part for part in [street_line, city_line] if part) or clean(row.get("address", ""))


def house_number_sort_key(value):
    match = re.match(r"(\d+)?(.*)", clean(value))
    number = int(match.group(1)) if match and match.group(1) else 0
    rest = match.group(2).lower() if match else ""
    return (number, rest)


def get_objects():
    with db() as conn:
        frame = pd.read_sql_query("SELECT * FROM objects", conn)
    if frame.empty:
        return frame
    sort_keys = frame.apply(
        lambda row: (clean(row["street"]).lower(), *house_number_sort_key(row["house_number"]), clean(row["name"]).lower()),
        axis=1,
    )
    return frame.assign(_sort_key=sort_keys).sort_values("_sort_key").drop(columns="_sort_key").reset_index(drop=True)


def get_assignments(object_id):
    with db() as conn:
        return pd.read_sql_query("""SELECT id, unit, location, usage_type, malo_id, meter_number, meter_location,
            measurement_location, tenant, status, checked_on, note, created_at, updated_at
            FROM assignments WHERE object_id = ? ORDER BY usage_type, unit""", conn, params=(object_id,))


def malo_valid(value):
    value = clean(value).replace(" ", "")
    return len(value) == 11 and value.isdigit()


def validation_report(frame):
    if frame.empty:
        return frame.assign(validation=[])
    report = frame.copy()
    report["malo_id"] = report["malo_id"].map(lambda x: clean(x).replace(" ", ""))
    report["meter_number"] = report["meter_number"].map(clean)
    malo_counts = report.loc[report["malo_id"] != "", "malo_id"].value_counts()
    meter_counts = report.loc[report["meter_number"] != "", "meter_number"].value_counts()
    issues_all = []
    for _, row in report.iterrows():
        issues = []
        malo, meter = row["malo_id"], row["meter_number"]
        if not clean(row["unit"]): issues.append("Wohnung/Einheit fehlt")
        if not malo: issues.append("MaLo-ID fehlt")
        elif not malo_valid(malo): issues.append("MaLo-ID muss 11 Ziffern haben")
        elif malo_counts.get(malo, 0) > 1: issues.append("MaLo-ID doppelt")
        if not meter: issues.append("Zählernummer fehlt")
        elif meter_counts.get(meter, 0) > 1: issues.append("Zählernummer doppelt")
        issues_all.append("; ".join(issues) if issues else "OK")
    report["validation"] = issues_all
    return report


def status_for_row(row):
    return "Prüfen" if row["validation"] != "OK" else (clean(row["status"]) or "Offen")


def add_assignment(conn, object_id, unit, location="", usage_type="Wohnung", malo_id="", meter_number="", meter_location="", measurement_location="", tenant="", status="Offen", checked_on="", note=""):
    now = datetime.now().isoformat(timespec="seconds")
    conn.execute("""INSERT INTO assignments (object_id, unit, location, usage_type, malo_id, meter_number,
        meter_location, measurement_location, tenant, status, checked_on, note, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (object_id, clean(unit), clean(location), clean(usage_type) or "Wohnung", clean(malo_id).replace(" ", ""),
         clean(meter_number), clean(meter_location), clean(measurement_location), clean(tenant), clean(status) or "Offen",
         clean(checked_on), clean(note), now, now))


def excel_export(object_row, frame):
    report = validation_report(frame)
    report["Exportstatus"] = report.apply(status_for_row, axis=1)
    wb = Workbook(); ws = wb.active; ws.title = "Zuordnung"
    headers = ["Wohnung / Einheit", "Name der Partei", "Lage", "Nutzungsart", "MaLo-ID", "Zählernummer", "Zählerplatz", "Messlokation", "Status", "Prüfdatum", "Bemerkung", "Validierung"]
    export_lines = [
        f"Straße / Hausnummer: {clean(object_row['street'])} {clean(object_row['house_number'])}".strip(),
        f"PLZ / Ort: {clean(object_row['postal_code'])} {clean(object_row['city'])}".strip(),
        f"Objektname: {object_row['name']}",
        "MaLo-ID / Stromzähler-Zuordnung",
        f"Exportiert: {datetime.now().strftime('%d.%m.%Y %H:%M')}", "", headers,
    ]
    for value in export_lines:
        ws.append(value if isinstance(value, list) else [value])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[7]: cell.font = Font(color="FFFFFF", bold=True); cell.fill = fill
    for _, row in report.iterrows():
        ws.append([clean(row["unit"]), clean(row["tenant"]), clean(row["location"]), clean(row["usage_type"]), clean(row["malo_id"]), clean(row["meter_number"]), clean(row["meter_location"]), clean(row["measurement_location"]), clean(row["Exportstatus"]), clean(row["checked_on"]), clean(row["note"]), clean(row["validation"])])
    for col in range(1, len(headers)+1):
        width = max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row+1)) + 2
        ws.column_dimensions[get_column_letter(col)].width = min(max(width, 14), 35)
    ws.freeze_panes = "A8"
    errors = wb.create_sheet("Prüfliste"); errors.append(["Wohnung / Einheit", "Name der Partei", "MaLo-ID", "Zählernummer", "Hinweis"])
    for cell in errors[1]: cell.font = Font(color="FFFFFF", bold=True); cell.fill = PatternFill("solid", fgColor="C00000")
    for _, row in report[report["validation"] != "OK"].iterrows(): errors.append([clean(row["unit"]), clean(row["tenant"]), clean(row["malo_id"]), clean(row["meter_number"]), clean(row["validation"])])
    for col in range(1, 6): errors.column_dimensions[get_column_letter(col)].width = 28
    details = wb.create_sheet("Objektdaten"); details.append(["Feld", "Wert"])
    for cell in details[1]: cell.font = Font(color="FFFFFF", bold=True); cell.fill = fill
    for label, value in [("Straße", object_row["street"]), ("Hausnummer", object_row["house_number"]), ("PLZ", object_row["postal_code"]), ("Ort", object_row["city"]), ("Objektname", object_row["name"]), ("Ansprechpartner", object_row["contact"]), ("Netzbetreiber", object_row["network_operator"])]: details.append([label, clean(value)])
    details.column_dimensions["A"].width = 24; details.column_dimensions["B"].width = 55
    output = io.BytesIO(); wb.save(output); output.seek(0); return output.getvalue()


def pdf_export(object_row, frame):
    report = validation_report(frame); report["Anzeige-Status"] = report.apply(status_for_row, axis=1)
    output = io.BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=1.0*cm, rightMargin=1.0*cm, topMargin=1.0*cm, bottomMargin=1.0*cm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"<b>{clean(object_row['street'])} {clean(object_row['house_number'])}</b>", styles["Title"]),
        Paragraph(f"{clean(object_row['postal_code'])} {clean(object_row['city'])}", styles["Heading2"]),
        Spacer(1, .18*cm),
        Paragraph(f"Objekt: {clean(object_row['name'])}", styles["Normal"]),
        Paragraph("MaLo-ID / Stromzähler-Zuordnung", styles["Heading2"]),
        Paragraph(f"Export: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles["Normal"]),
        Spacer(1, .3*cm),
    ]
    data = [["Wohnung", "Name der Partei", "Lage", "MaLo-ID", "Zählernummer", "Zählerplatz", "Status", "Hinweis"]]
    for _, row in report.iterrows(): data.append([clean(row["unit"]), clean(row["tenant"]), clean(row["location"]), clean(row["malo_id"]), clean(row["meter_number"]), clean(row["meter_location"]), clean(row["Anzeige-Status"]), "" if row["validation"] == "OK" else clean(row["validation"])])
    table = Table(data, colWidths=[2.0*cm, 3.5*cm, 2.2*cm, 3.0*cm, 3.0*cm, 3.0*cm, 2.0*cm, 5.1*cm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),7.2),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#B7C9D6")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#EAF2F8")]),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
    story.extend([table, Spacer(1, .55*cm), Paragraph("Prüfung / Abnahme: ____________________________________    Datum: ______________", styles["Normal"])])
    doc.build(story); output.seek(0); return output.getvalue()


def import_excel(uploaded, object_id):
    frame = pd.read_excel(uploaded); frame.columns = [str(c).strip().lower() for c in frame.columns]
    aliases = {"einheit":"unit","wohnung":"unit","top":"unit","name":"tenant","name der partei":"tenant","partei":"tenant","mieter":"tenant","mieter/kunde":"tenant","lage":"location","nutzungsart":"usage_type","malo-id":"malo_id","malo id":"malo_id","zählernummer":"meter_number","zaehlernummer":"meter_number","zählerplatz":"meter_location","zaehlerplatz":"meter_location","messlokation":"measurement_location","status":"status","prüfdatum":"checked_on","pruefdatum":"checked_on","bemerkung":"note"}
    frame = frame.rename(columns={c: aliases.get(c, c) for c in frame.columns})
    if frame.columns.duplicated().any():
        frame = frame.loc[:, ~frame.columns.duplicated()]
    if "unit" not in frame.columns: raise ValueError("In der Importdatei fehlt die Spalte 'Wohnung', 'Einheit' oder 'Top'.")
    with db() as conn:
        for _, row in frame.iterrows():
            values = [clean(row[field]) if field in frame.columns else ("Wohnung" if field == "usage_type" else "Offen" if field == "status" else "") for field in ["unit","location","usage_type","malo_id","meter_number","meter_location","measurement_location","tenant","status","checked_on","note"]]
            add_assignment(conn, object_id, *values)


init_db()
st.title("⚡ MaLo-ID & Stromzähler-Zuordnung")
st.caption("Für Mehrfamilienhäuser mit bis zu 15 Parteien je Haus – inklusive Namen, MaLo-ID, Zählernummer sowie Excel- und PDF-Export.")

with st.sidebar:
    st.header("Objektverwaltung")
    objects = get_objects(); options = {f"{object_address(row)} – {row['name']}": int(row["id"]) for _, row in objects.iterrows()}
    selected_label = st.selectbox("Aktives Objekt", ["— Neues Objekt anlegen —", *options.keys()])
    st.divider(); st.subheader("Neues Haus anlegen")
    with st.form("new_object", clear_on_submit=True):
        a,b = st.columns([3,1]); new_street = a.text_input("Straße *", placeholder="Musterstraße"); new_house_number = b.text_input("Hausnummer *", placeholder="12A")
        c,d = st.columns([1,3]); new_postal_code = c.text_input("PLZ", placeholder="1010"); new_city = d.text_input("Ort", placeholder="Wien")
        new_name = st.text_input("Objektname", placeholder="z. B. MFH Musterstraße 12")
        new_contact = st.text_input("Ansprechpartner"); new_operator = st.text_input("Netzbetreiber")
        create_object = st.form_submit_button("Haus anlegen", type="primary")
    if create_object:
        if not new_street.strip() or not new_house_number.strip(): st.error("Bitte Straße und Hausnummer eingeben.")
        else:
            display_name = new_name.strip() or f"MFH {new_street.strip()} {new_house_number.strip()}"
            with db() as conn: conn.execute("INSERT INTO objects (name, street, house_number, postal_code, city, contact, network_operator, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (display_name,new_street.strip(),new_house_number.strip(),new_postal_code.strip(),new_city.strip(),new_contact.strip(),new_operator.strip(),datetime.now().isoformat(timespec="seconds")))
            st.success("Haus angelegt."); st.rerun()

if selected_label == "— Neues Objekt anlegen —":
    st.info("Lege links zuerst ein Haus an. Straße und Hausnummer stehen dabei ganz oben.")
    st.markdown("""### Schneller Ablauf für bis zu 15 Parteien
1. Ganz oben Straße und Hausnummer des Objekts eingeben.
2. Haus anlegen.
3. Im Tab **15 Parteien anlegen** Wohnungen und Namen eintragen.
4. MaLo-ID, Zählernummer und Zählerplatz ergänzen.
5. Als Excel oder PDF exportieren.
""")
    st.stop()

object_id = options[selected_label]
with db() as conn: object_row = pd.read_sql_query("SELECT * FROM objects WHERE id = ?", conn, params=(object_id,)).iloc[0]
assignments = get_assignments(object_id); report = validation_report(assignments)

st.markdown(f"# {clean(object_row['street'])} {clean(object_row['house_number'])}")
st.markdown(f"### {clean(object_row['postal_code'])} {clean(object_row['city'])}")
st.caption(f"Objekt: {clean(object_row['name'])}")
c1,c2,c3,c4 = st.columns(4); c1.metric("Einträge",len(report)); c2.metric("Wohnungen",int((report["usage_type"]=="Wohnung").sum()) if not report.empty else 0); c3.metric("Offene Hinweise",int((report["validation"]!="OK").sum()) if not report.empty else 0); c4.metric("MaLo-IDs vorhanden",int((report["malo_id"].map(clean)!="").sum()) if not report.empty else 0)

tab_batch, tab_add, tab_list, tab_export, tab_import = st.tabs(["15 Parteien anlegen", "Eintrag ergänzen", "Zuordnungsliste", "Export", "Excel-Import"])

with tab_batch:
    st.markdown("### Wohnungen und Namen für bis zu 15 Parteien anlegen")
    st.write("Trage pro Zeile mindestens die Wohnung/Einheit und bei Bedarf den Namen der Partei ein. Leere Zeilen werden nicht gespeichert.")
    defaults = pd.DataFrame({"Wohnung / Einheit": [f"Wohnung {i:02d}" for i in range(1,16)], "Name der Partei": [""]*15, "Lage": [""]*15})
    with st.form("batch_parties"):
        party_frame = st.data_editor(defaults, num_rows="fixed", hide_index=True, use_container_width=True)
        save_parties = st.form_submit_button("Parteien als Wohnungen anlegen", type="primary")
    if save_parties:
        valid_rows = party_frame[party_frame["Wohnung / Einheit"].fillna("").astype(str).str.strip() != ""]
        if valid_rows.empty: st.warning("Bitte mindestens eine Wohnung eintragen.")
        else:
            with db() as conn:
                for _, row in valid_rows.iterrows(): add_assignment(conn, object_id, row["Wohnung / Einheit"], row["Lage"], "Wohnung", tenant=row["Name der Partei"], checked_on=date.today().strftime("%d.%m.%Y"))
            st.success(f"{len(valid_rows)} Wohnungen/Parteien angelegt. Ergänze nun MaLo-ID und Zählernummer."); st.rerun()

with tab_add:
    st.markdown("### Einzelne Zuordnung ergänzen")
    with st.form("add_assignment", clear_on_submit=True):
        a,b,c = st.columns(3); unit=a.text_input("Wohnung / Einheit *",placeholder="z. B. Wohnung 07"); tenant=b.text_input("Name der Partei",placeholder="z. B. Maria Muster"); location=c.text_input("Lage",placeholder="z. B. 2. OG links")
        d,e,f = st.columns(3); usage_type=d.selectbox("Nutzungsart",["Wohnung","Allgemeinstrom","Gewerbe","Heizung","Aufzug","Tiefgarage","PV-Anlage","Sonstiges"]); malo_id=e.text_input("MaLo-ID",max_chars=11,placeholder="11 Ziffern"); meter_number=f.text_input("Zählernummer",placeholder="z. B. EHZ123456789")
        g,h,i = st.columns(3); meter_location=g.text_input("Zählerplatz"); measurement_location=h.text_input("Messlokation (optional)"); status=i.selectbox("Status",["Offen","Geprüft","Unklar"])
        checked_on=st.date_input("Prüfdatum",value=date.today(),format="DD.MM.YYYY"); note=st.text_area("Bemerkung")
        save=st.form_submit_button("Zuordnung speichern",type="primary")
    if save:
        if not unit.strip(): st.error("Wohnung / Einheit ist ein Pflichtfeld.")
        else:
            with db() as conn: add_assignment(conn,object_id,unit,location,usage_type,malo_id,meter_number,meter_location,measurement_location,tenant,status,checked_on.strftime("%d.%m.%Y"),note)
            st.success("Zuordnung gespeichert."); st.rerun()

with tab_list:
    st.markdown("### Vollständige Zuordnung")
    if report.empty: st.info("Für dieses Haus sind noch keine Parteien oder Zuordnungen erfasst.")
    else:
        filter_text=st.text_input("Suche",placeholder="Wohnung, Name, MaLo-ID, Zählernummer oder Zählerplatz")
        shown=report.copy()
        if filter_text.strip():
            cols=["unit","tenant","location","malo_id","meter_number","meter_location"]
            shown=shown[shown[cols].fillna("").astype(str).apply(lambda col: col.str.contains(filter_text,case=False,na=False)).any(axis=1)]
        editable_cols = ["unit","tenant","location","usage_type","malo_id","meter_number","meter_location","status","checked_on","note"]
        column_labels = {"id":"ID","unit":"Wohnung / Einheit","tenant":"Name der Partei","location":"Lage","usage_type":"Nutzungsart","malo_id":"MaLo-ID","meter_number":"Zählernummer","meter_location":"Zählerplatz","status":"Status","checked_on":"Prüfdatum","note":"Bemerkung","validation":"Prüfung"}
        display=shown[["id",*editable_cols,"validation"]].rename(columns=column_labels)
        st.caption("Werte direkt in der Tabelle bearbeiten und anschließend über den Button unten speichern.")
        edited=st.data_editor(
            display,use_container_width=True,hide_index=True,disabled=["ID","Prüfung"],key="assignments_editor",
            column_config={
                "Nutzungsart": st.column_config.SelectboxColumn(options=["Wohnung","Allgemeinstrom","Gewerbe","Heizung","Aufzug","Tiefgarage","PV-Anlage","Sonstiges"]),
                "Status": st.column_config.SelectboxColumn(options=["Offen","Geprüft","Unklar"]),
            },
        )
        if st.button("Änderungen speichern",type="primary"):
            reverse_labels={v:k for k,v in column_labels.items()}
            edited_internal=edited.rename(columns=reverse_labels)
            original_by_id=shown.set_index("id")
            updated=0
            with db() as conn:
                for _,row in edited_internal.iterrows():
                    row_id=int(row["id"])
                    if row_id not in original_by_id.index: continue
                    original=original_by_id.loc[row_id]
                    new_values={col: clean(row[col]) for col in editable_cols}
                    if any(new_values[col] != clean(original[col]) for col in editable_cols):
                        assignments_set=", ".join(f"{col} = ?" for col in editable_cols)
                        conn.execute(f"UPDATE assignments SET {assignments_set}, updated_at = ? WHERE id = ?",(*[new_values[col] for col in editable_cols],datetime.now().isoformat(timespec="seconds"),row_id))
                        updated+=1
            if updated: st.success(f"{updated} Eintrag/Einträge aktualisiert."); st.rerun()
            else: st.info("Keine Änderungen erkannt.")
        st.markdown("### Eintrag löschen")
        delete_options={f"{int(row['id'])} – {row['unit']} / {clean(row['tenant'])}":int(row["id"]) for _,row in shown.iterrows()}
        if delete_options:
            delete_label=st.selectbox("Eintrag auswählen",list(delete_options.keys()))
            if st.button("Ausgewählten Eintrag löschen"):
                with db() as conn: conn.execute("DELETE FROM assignments WHERE id = ?",(delete_options[delete_label],))
                st.success("Eintrag gelöscht."); st.rerun()

with tab_export:
    st.markdown("### Excel und PDF exportieren")
    if report.empty: st.warning("Bitte zuerst Parteien oder Zuordnungen erfassen.")
    else:
        errors=report[report["validation"]!="OK"]
        if not errors.empty: st.warning(f"Es gibt {len(errors)} Einträge mit Hinweisen. Diese werden im Export gekennzeichnet.")
        safe_name="".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in f"{object_row['street']}_{object_row['house_number']}").strip("_") or "haus"
        left,right=st.columns(2)
        left.download_button("Excel-Zuordnungsliste herunterladen",data=excel_export(object_row,assignments),file_name=f"MaLo_Zuordnung_{safe_name}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary")
        right.download_button("PDF-Zuordnungsliste herunterladen",data=pdf_export(object_row,assignments),file_name=f"MaLo_Zuordnung_{safe_name}.pdf",mime="application/pdf",type="primary")

with tab_import:
    st.markdown("### Bestehende Excel-Liste importieren")
    st.write("Unterstützte Spalten: Wohnung/Einheit/Top, Name/Name der Partei/Mieter, Lage, MaLo-ID, Zählernummer, Zählerplatz, Messlokation, Status, Prüfdatum und Bemerkung.")
    uploaded=st.file_uploader("Excel-Datei auswählen",type=["xlsx","xls"])
    if uploaded is not None and st.button("Datei in dieses Haus importieren",type="primary"):
        try:
            import_excel(uploaded,object_id); st.success("Import abgeschlossen. Bitte prüfe anschließend die Zuordnungsliste."); st.rerun()
        except Exception as exc: st.error(f"Import nicht möglich: {exc}")

st.divider(); st.caption("Hinweis: Die App prüft Struktur und Dubletten. Die fachliche Richtigkeit der Zuordnung ist anhand von Netzbetreiber-, Zähler- und Vertragsunterlagen zu bestätigen.")
